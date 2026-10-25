# Lap danh sach DE XUAT doi ten may Samsung con ton: them mau tieng Anh (mau in tren hop) + ma may.
# Chi doc, khong ghi DB. Ket qua la Excel de chu shop duyet; phien PC moi sua trong HTSoft sau khi duyet.
# Quy tac ten: [Loai] SAMSUNG [Model] [RAM] [ROM] [Mau VN] [Mau EN] [Ma may]
#   vd: DTDD SAMSUNG A07 5G 4GB 128GB Tim Light Violet SM-A076BLVBXXV
# Mau EN lay tu 2 ky tu ngay sau model trong ma may: Z = mau co ban, L = Light, D = Dark, A = Absolute.
# Cach chay: python doi_ten_samsung.py <ban-chup.json> <ra.xlsx>
#   ban-chup.json: ban chup app Deal Samsung (co mang 'hang' va 'map'), hoac file chi co mang 'hang'.
import json, re, sys
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Tu dien goc, da khop 100% voi bang map trong HTSoft (321 ma, 15/10/2026).
MA_MAU = {
    'ZK': 'Black', 'ZW': 'White', 'ZS': 'Silver', 'ZA': 'Gray', 'ZB': 'Blue', 'ZG': 'Green',
    'ZV': 'Violet', 'ZR': 'Red', 'ZD': 'Gold', 'ZE': 'Beige', 'ZT': 'Titanium',
    'LB': 'Light Blue', 'LG': 'Light Green', 'LV': 'Light Violet', 'LI': 'Light Pink',
    'DB': 'Dark Blue', 'DG': 'Dark Green', 'DA': 'Dark Gray',
    'AK': 'Absolute Black', 'AW': 'Absolute White', 'ID': 'Pink Gold',
}
# Ma suy luan, chua co trong map -> luon 'can xac nhan' (doc mau tren hop).
MA_MAU_SUY = {'ZY': 'Yellow'}
# Ma day du ket thuc bang 1 ky tu phien ban + 3 ky tu vung (XXV, INS...): mau = 2 ky tu ngay truoc 4 ky tu cuoi.
#   SM-A076BLVBXXV -> mau LV. Ma cut (SM-A075F, SM-M526) khong co mau.
RE_MA = re.compile(r'\bSM-[A-Z]\d{3}[A-Z0-9]{3,6}[A-Z]{3}\b', re.I)
RE_MA_BAT_KY = re.compile(r'\bSM-[A-Z0-9]+\b', re.I)
RE_BO_NHO = re.compile(r'^\d+(GB|G|TB)$', re.I)
# Mau VN goc -> cac mau EN hop le, de bat ten ghi lech mau voi ma.
VN_EN = {
    'ĐEN': {'Black', 'Absolute Black', 'Titanium'}, 'TRẮNG': {'White', 'Absolute White'},
    'BẠC': {'Silver', 'Gray', 'Light Violet', 'Titanium'}, 'XÁM': {'Gray', 'Dark Gray', 'Silver', 'Titanium'},
    'XANH': {'Blue', 'Light Blue', 'Dark Blue', 'Green', 'Light Green', 'Dark Green'},
    'TÍM': {'Violet', 'Light Violet'}, 'VÀNG': {'Gold', 'Yellow', 'Pink Gold'}, 'HỒNG': {'Light Pink', 'Pink Gold'},
    'ĐỎ': {'Red'}, 'KEM': {'Beige', 'White', 'Light Green'}, 'BE': {'Beige'}, 'MINT': {'Light Green', 'Green'},
}


def tach_ma(chuoi):
    """Tra ve (ma day du, ma mau 2 ky tu) hoac (ma cut, None) hoac (None, None)."""
    for m in RE_MA.finditer(chuoi or ''):
        ma = m.group(0).upper()
        if ma[-6:-4].isalpha():
            return ma, ma[-6:-4]
    m = RE_MA_BAT_KY.search(chuoi or '')
    return (m.group(0).upper(), None) if m else (None, None)


def tach_ten(ten, bo=()):
    """(phan dau den het bo nho, phan mau) - bo ma SM-... va cac tu trong 'bo' (tu cua mau EN se them, tranh lap)."""
    tu = [w for w in ten.split() if not RE_MA_BAT_KY.fullmatch(w)]
    vt = max((i for i, w in enumerate(tu) if RE_BO_NHO.match(w)), default=-1)
    if vt < 0:
        return None, None
    mau = [w for w in tu[vt + 1:] if w.upper() not in bo]
    return tu[:vt + 1], mau


def mau_goc(mau_vn):
    return [w.upper() for w in mau_vn if w.upper() in VN_EN]


def lap(hang, ban_do):
    theo_mhid = {m['MHID']: m for m in ban_do if m.get('ConHieuLuc', True)}
    ra = {'ngay': [], 'xac_nhan': [], 'thieu_ma': []}
    for x in hang:
        if not (x.get('Hang') or '').upper().startswith('SAMSUNG') or (x.get('Ton') or 0) <= 0:
            continue
        ten = x['MHTEN'].strip()
        ma, mau = tach_ma(x.get('SKU'))
        if not mau:
            ma2, mau2 = tach_ma(ten)
            if mau2 or not ma:
                ma, mau = ma2, mau2
        if not mau and x.get('MHID') in theo_mhid:
            ma, mau = tach_ma(theo_mhid[x['MHID']]['SKU'])
        en = MA_MAU.get(mau or '') or MA_MAU_SUY.get(mau or '', '')
        dong = {'MHID': x.get('MHID'), 'MHCODE': x.get('MHCODE'), 'TEN_CU': ten, 'TON': x.get('Ton'),
                'MA_MAY': ma or '', 'MA_MAU': mau or '', 'MAU_EN': en, 'TEN_MOI': '', 'GHI_CHU': ''}
        if not mau:
            dong['GHI_CHU'] = (f'Mã máy bị cắt ({ma}), thiếu mã màu' if ma else 'HTSoft chưa có mã máy SM-…') + \
                              ' → xem hộp ghi mã đầy đủ'
            ra['thieu_ma'].append(dong); continue
        if not en:
            dong['GHI_CHU'] = f'Mã màu {mau} chưa có trong từ điển → đọc màu trên hộp'
            ra['xac_nhan'].append(dong); continue
        bo = {w.upper() for w in en.split()}
        dau, vn = tach_ten(ten, bo)
        _, vn_mh = tach_ten(x.get('MHCODE') or '', bo)
        ghi = []
        if dau is None:
            dong['GHI_CHU'] = 'Không tách được RAM/ROM trong tên → sửa tay'
            ra['xac_nhan'].append(dong); continue
        # Ten bi don nhieu mau (vd 'TIM NHAT DEN XANH LA') ma Ma hang giu 1 mau dung -> lay mau theo Ma hang.
        if vn_mh and [w.upper() for w in vn] != [w.upper() for w in vn_mh] and \
                ' '.join(vn).upper().endswith(' '.join(vn_mh).upper()):
            ghi.append('Tên cũ ghi dồn nhiều màu (' + ' '.join(vn) + ') → giữ màu theo Mã hàng')
            vn = vn_mh
        if not vn:
            ghi.append('Tên chưa có màu → thêm màu VN trước ' + en + '; nếu đây là mã cha (gộp nhiều màu) thì KHÔNG đổi')
        else:
            goc = mau_goc(vn)
            if goc and en not in VN_EN[goc[0]]:
                ghi.append(f'Màu VN "{" ".join(vn)}" lệch mã màu {mau}={en} → kiểm hộp')
            elif len(set(goc)) >= 3:
                ghi.append('Tên còn ' + str(len(set(goc))) + ' màu → giữ 1 màu đúng với mã')
        if mau in MA_MAU_SUY:
            ghi.append(f'{mau}={en} là suy luận, chưa có trong bảng map → đối chiếu hộp')
        dong['TEN_MOI'] = ' '.join(dau + vn + en.split() + [ma])
        dong['GHI_CHU'] = '; '.join(ghi)
        ra['xac_nhan' if ghi else 'ngay'].append(dong)
    return ra


COT = [('MHCODE', 'Mã hàng', 34), ('TEN_CU', 'Tên cũ', 58), ('TEN_MOI', 'Tên mới đề xuất', 62),
       ('MA_MAY', 'Mã máy', 18), ('MA_MAU', 'Mã màu', 8), ('MAU_EN', 'Màu EN (hộp)', 16),
       ('TON', 'Tồn', 6), ('GHI_CHU', 'Ghi chú', 50), ('MHID', 'MHID (cho phiên PC)', 38)]


def ghi_excel(ra, duong_dan, luc):
    wb = Workbook(); wb.remove(wb.active)
    to = PatternFill('solid', fgColor='DDEBF7')
    for ten_sheet, khoa in [('1 Doi duoc ngay', 'ngay'), ('2 Can xac nhan', 'xac_nhan'), ('3 Thieu ma may', 'thieu_ma')]:
        ws = wb.create_sheet(ten_sheet)
        ws.append([c[1] for c in COT])
        for i, c in enumerate(COT, 1):
            ws.cell(1, i).font = Font(bold=True); ws.cell(1, i).fill = to
            ws.column_dimensions[get_column_letter(i)].width = c[2]
        for d in sorted(ra[khoa], key=lambda d: d['TEN_CU']):
            ws.append([d[c[0]] for c in COT])
        ws.freeze_panes = 'A2'
    ws = wb.create_sheet('4 Tu dien ma mau')
    ws.append(['Mã màu', 'Màu EN', 'Quy tắc'])
    for k, v in list(MA_MAU.items()) + list(MA_MAU_SUY.items()):
        ws.append([k, v, {'Z': 'màu cơ bản', 'L': 'Light = nhạt', 'D': 'Dark = đậm', 'A': 'Absolute'}.get(k[0], 'riêng') + (' — suy luận, cần xác nhận' if k in MA_MAU_SUY else '')])
    ws.append([]); ws.append([f'Nguồn tồn: bản chụp HTSoft lúc {luc}. Phiên PC kiểm lại tồn trước khi sửa.'])
    wb.save(duong_dan)


if __name__ == '__main__':
    du_lieu = json.load(open(sys.argv[1], encoding='utf-8'))
    ra = lap(du_lieu['hang'], du_lieu.get('map', []))
    ghi_excel(ra, sys.argv[2], du_lieu.get('luc', '?'))
    print({k: len(v) for k, v in ra.items()})
